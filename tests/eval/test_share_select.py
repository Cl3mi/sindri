import json
import sys
from pathlib import Path

import fitz
import openpyxl
import pytest

from app.eval.anon import Anonymizer
from app.eval.dump import save_dump
from app.eval.models import PredictionDump, RunConfig
from app.eval.share_select import main as select_main
from app.eval.share_select import manifest, select
from app.models import Characteristic, ExtractionResult

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import export_share  # noqa: E402


def _row(doc, recall, n_gold=20, n_pred=25, escaped=0, missed=2, false=5, run="r"):
    return {"doc": doc, "run": run, "n_gold": n_gold, "n_pred": n_pred,
            "recall": recall,
            "counts": {"escaped_error": escaped, "missed": missed,
                       "false_detection": false}}


def test_great_are_best_recall_first_and_roles_are_tagged():
    rows = [_row("a", .9), _row("b", .7), _row("c", .6), _row("d", .56)]
    picks = select(rows, great=3, problem=0)
    assert [p["doc"] for p in picks] == ["a", "b", "c"]
    assert {p["role"] for p in picks} == {"great"}


def test_small_and_escape_prone_drawings_are_never_picked():
    rows = [_row("tiny", .9, n_gold=5), _row("leaky", .9, escaped=4),
            _row("fine", .6)]
    assert [p["doc"] for p in select(rows, 3, 2)] == ["fine"]


def test_problem_picks_are_mid_recall_one_per_flavour():
    rows = [_row("great", .8),
            _row("over", .45, n_pred=60, false=40, missed=1),
            _row("miss", .4, n_pred=20, false=2, missed=9),
            _row("broken", .1)]
    picks = select(rows, great=1, problem=2)
    assert [(p["doc"], p["role"]) for p in picks] == [
        ("great", "great"), ("over", "problem"), ("miss", "problem")]
    assert "over-detection" in picks[1]["why"] and "misses" in picks[2]["why"]


def test_short_pool_returns_fewer_without_crashing():
    assert select([_row("a", .9)], great=3, problem=2) == select(
        [_row("a", .9)], great=1, problem=0)


def test_manifest_groups_hashes_by_run():
    picks = select([_row("a", .9, run="dev"), _row("b", .8, run="test")], 2, 0)
    assert manifest(picks)["runs"] == [{"run": "dev", "hashes": ["a"]},
                                       {"run": "test", "hashes": ["b"]}]


def test_train_digest_is_refused(tmp_path):
    f = tmp_path / "t.json"
    f.write_text(json.dumps({"split": "train", "per_doc": []}))
    with pytest.raises(SystemExit):
        select_main(["--digest", f"r={f}", "--out", str(tmp_path / "m.json")])


def test_train_digest_is_allowed_only_when_asked(tmp_path):
    """The review campaign selects from train + dev (plan C1 decision 5): 50
    drawings exceed dev's scoped 15. A client share-pack still refuses train
    unless the flag is given."""
    f = tmp_path / "t.json"
    f.write_text(json.dumps({"split": "train", "per_doc": [_row("a", .9)]}))
    out = tmp_path / "m.json"
    assert select_main(["--digest", f"r={f}", "--out", str(out), "--great", "1",
                        "--problem", "0", "--allow-train"]) == 0
    assert json.loads(out.read_text())["roles"] == {"a": "great"}


@pytest.mark.parametrize("split", ["test", "all", "", None])
def test_the_campaign_never_draws_from_test(tmp_path, split):
    """--allow-train is the review campaign's flag, and the campaign keeps
    test held out (decision 5): the client's grades must not have touched
    the split every later arm is confirmed on. "all" (an unsplit score)
    contains test drawings, and a digest that does not say its split cannot
    prove it holds none, so under --allow-train only train and dev pass."""
    f = tmp_path / "t.json"
    digest = {"per_doc": [_row("a", .9)]}
    if split is not None:
        digest["split"] = split
    f.write_text(json.dumps(digest))
    with pytest.raises(SystemExit):
        select_main(["--digest", f"r={f}", "--out", str(tmp_path / "m.json"),
                     "--allow-train"])
    assert not (tmp_path / "m.json").exists()


def test_a_dev_digest_still_passes_with_allow_train(tmp_path):
    f = tmp_path / "d.json"
    f.write_text(json.dumps({"split": "dev", "per_doc": [_row("a", .9)]}))
    assert select_main(["--digest", f"r={f}", "--out", str(tmp_path / "m.json"),
                        "--great", "1", "--problem", "0", "--allow-train"]) == 0


def test_end_to_end_manifest_to_files(tmp_path, monkeypatch):
    monkeypatch.setenv("SINDRI_DOC_SALT", "testsalt")
    monkeypatch.delenv("CLAUDECODE", raising=False)   # main() refuses under an agent
    root = tmp_path / "root"
    (root / "runs" / "r1").mkdir(parents=True)
    (root / "corpus" / "originals").mkdir(parents=True)
    s = 300 / 72
    doc = fitz.open()
    doc.new_page(width=600, height=400)
    doc.save(root / "corpus" / "originals" / "PART-1.pdf")
    rows = [Characteristic(pos=1, nominal="10", balloon_xy=(100 * s, 100 * s),
                           target_region=(200 * s, 150 * s, 260 * s, 170 * s)),
            Characteristic(pos=2, nominal="9", suggested=True)]
    save_dump(PredictionDump(doc_id="PART-1", config=RunConfig(), scale=s,
                             page_rect=(0, 0, 600, 400),
                             result=ExtractionResult(characteristics=rows,
                                                     render_scale=s)),
              root / "runs" / "r1")
    h = Anonymizer()("PART-1")
    man = tmp_path / "m.json"
    man.write_text(json.dumps({"runs": [{"run": "r1", "hashes": [h]}]}))
    out = tmp_path / "out"

    assert export_share.main(["--root", str(root), "--manifest", str(man),
                              "--out", str(out)]) == 0
    ws = openpyxl.load_workbook(out / "PART-1.xlsx").active
    assert [ws.cell(3, 1).value, ws.cell(3, 3).value] == [1, "10"]
    assert ws.max_row == 3                      # the suggestion was not exported
    assert len(fitz.open(out / "PART-1_ballooned.pdf")[0].get_drawings()) == 2


def test_unknown_hash_fails_loudly(tmp_path, monkeypatch):
    monkeypatch.setenv("SINDRI_DOC_SALT", "testsalt")
    monkeypatch.delenv("CLAUDECODE", raising=False)   # main() refuses under an agent
    # The out dir sits beside the root, never inside it (check_out_dir).
    (tmp_path / "root" / "runs" / "r1").mkdir(parents=True)
    man = tmp_path / "m.json"
    man.write_text(json.dumps({"runs": [{"run": "r1", "hashes": ["deadbeef"]}]}))
    assert export_share.main(["--root", str(tmp_path / "root"), "--manifest", str(man),
                              "--out", str(tmp_path / "o")]) == 1


def test_manifest_records_each_picks_role():
    """The review campaign stores great/problem per drawing, so its report
    can reweight the 60/40 pick (campaign design §5). Picks carry the drawing
    id under "doc" (what `take()` writes); the manifest calls it a hash."""
    from app.eval.share_select import manifest
    picks = [{"run": "r5-control", "doc": "aa", "role": "great"},
             {"run": "r3-trainpredict", "doc": "bb", "role": "problem"}]
    m = manifest(picks)
    assert m["roles"] == {"aa": "great", "bb": "problem"}
    assert {r["run"] for r in m["runs"]} == {"r5-control", "r3-trainpredict"}


def _pool_rows():
    return [{"run": "r5-control", "doc": f"{i:08x}", "n_gold": 20, "n_pred": 20,
             "recall": rec, "counts": {"escaped_error": 0}}
            for i, rec in enumerate([0.9, 0.8, 0.6, 0.4, 0.35, 0.1])]


def test_manifest_records_the_eligible_pool_for_reweighting():
    """A 3+2 pick over a 3-great / 2-problem pool says nothing about the
    population until the grader knows how big each eligible set was, so the
    manifest carries the counts of the ELIGIBLE sets (not of what was left
    after picking) for campaign_grade to reweight by."""
    from app.eval.share_select import eligible_counts, manifest
    pool = eligible_counts(_pool_rows())
    assert pool == {"great": 3, "problem": 2}
    assert manifest([], pool)["pool"] == pool
    assert "pool" not in manifest([])           # optional: old manifests unchanged


def test_eligible_counts_and_select_share_one_filter():
    """They must never drift: the pool is reweighting's denominator, and a
    filter select() applied but the count did not would silently skew it."""
    from app.eval.share_select import eligible_counts, select
    rows = _pool_rows() + [
        {"run": "r5-control", "doc": "ffffffff", "n_gold": 5, "n_pred": 5,
         "recall": 0.9, "counts": {}},                       # too few gold rows
        {"run": "r5-control", "doc": "eeeeeeee", "n_gold": 20, "n_pred": 20,
         "recall": 0.9, "counts": {"escaped_error": 10}},    # silent errors
        {"run": "r5-control", "doc": "dddddddd", "n_gold": 20, "n_pred": 20,
         "recall": None, "counts": {}}]                      # unmeasured
    assert eligible_counts(rows) == {"great": 3, "problem": 2}
    picks = select(rows, great=99, problem=99)
    assert sum(p["role"] == "great" for p in picks) == 3
    # a looser escaped ceiling admits the noisy drawing in both
    assert eligible_counts(rows, max_escaped=0.6)["great"] == 4


def test_manifest_with_pool_passes_the_grader():
    """Contract test: the grader refuses a malformed roles/pool/runs, and the
    manifest this module writes is its only producer."""
    from app.eval.campaign_grade import _manifest_parts
    from app.eval.share_select import eligible_counts, manifest, select
    rows = _pool_rows()
    m = manifest(select(rows, 2, 1), eligible_counts(rows))
    roles, pool, runs = _manifest_parts(json.loads(json.dumps(m)))
    assert pool == {"great": 3, "problem": 2}
    assert runs == {"r5-control"} and set(roles.values()) <= {"great", "problem"}
