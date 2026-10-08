import hashlib
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from app.main import app
from app.pipeline.detect import Detection
from tests.conftest import StubVLMBackend

client = TestClient(app)


def save_pdf(test_client, path, filename="sample.pdf"):
    """Import a PDF the way the UI does -- a draft batch of one, the chunked
    upload protocol, completion -- and return its metadata. Nothing is
    queued."""
    data = Path(path).read_bytes()
    batch = test_client.post("/api/batches", json={"name": filename}).json()
    r = test_client.post("/api/uploads", json={
        "batch_id": batch["id"], "name": filename, "size": len(data),
        "fingerprint": hashlib.sha256(data).hexdigest()})
    assert r.status_code == 200, r.text
    uid = r.json()["upload_id"]
    if data:
        assert test_client.put(f"/api/uploads/{uid}?offset=0", content=data).status_code == 200
    r = test_client.post(f"/api/uploads/{uid}/complete")
    assert r.status_code == 200, r.text
    (row,) = r.json()["documents"]
    return {"session_id": row["id"], "fileName": row["file_name"],
            "pages": row["pages"], "batch_id": batch["id"]}


def start_pdf(test_client, sid):
    """What the interim UI's "Include and start" does for its batch of one:
    include the drawing, then start the import -- naming the drawing it
    expects the draft to hold, as the UI does."""
    batch_id = test_client.get(f"/api/documents/{sid}").json()["batch_id"]
    r = test_client.post(f"/api/documents/{sid}/include", json={"included": True})
    assert r.status_code == 200, r.text
    r = test_client.post(f"/api/batches/{batch_id}/start", json={"expected_ids": [sid]})
    assert r.status_code == 200, r.text
    assert r.json()["documents"] == [sid]


def upload_pdf(test_client, path, filename="sample.pdf"):
    """Import, queue and extract a PDF through the background worker -- run
    inline, so the test is deterministic -- then open it for review. Returns
    the /open payload (session_id, image_url, rows, notes, marks,
    title_block, writer, events, last_seq)."""
    import app.main as main
    meta = save_pdf(test_client, path, filename)
    sid = meta["session_id"]
    start_pdf(test_client, sid)
    assert main.runtime().worker.run_once() == sid
    doc = test_client.get(f"/api/documents/{sid}").json()
    assert doc["state"] == "ready", doc
    r = test_client.post(f"/api/documents/{sid}/open")
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture
def stub_backend(monkeypatch):
    backend = StubVLMBackend(
        detections=[Detection((40, 40, 120, 70), "dimension", 0.9)],
        text="1,2 +0,1 -0,1",
    )
    monkeypatch.setattr("app.main._BACKEND", backend)
    return backend


def test_the_single_shot_upload_and_per_drawing_enqueue_are_gone():
    """Since B2 a drawing arrives through /api/uploads into a draft and is
    queued only by starting the draft, where the Include rules live. Asked
    of the route table: a request to an unknown id at /enqueue would 404
    whether the route existed or not."""
    routes = {(m, r.path) for r in app.routes for m in (getattr(r, "methods", None) or ())}
    assert ("POST", "/api/upload") not in routes
    assert ("POST", "/api/documents/{doc_id}/enqueue") not in routes
    assert ("POST", "/api/batches/{batch_id}/start") in routes


def test_every_test_gets_its_own_data_and_review_dir(tmp_path):
    import app.main as main
    rt = main.runtime()
    assert rt.data.root == tmp_path / "data"
    assert rt.store.root == tmp_path / "data" / "reviews"


def test_upload_returns_rows_and_image(sample_pdf, stub_backend):
    data = upload_pdf(client, sample_pdf)
    assert "session_id" in data
    assert len(data["rows"]) >= 1
    assert data["rows"][0]["source"] == "auto"
    assert data["rows"][0]["id"]
    assert data["image_url"].startswith("/api/image/")


def test_queued_drawings_report_their_queue_position(sample_pdf, stub_backend):
    a = save_pdf(client, sample_pdf)["session_id"]
    b = save_pdf(client, sample_pdf)["session_id"]     # a duplicate: start_pdf includes it
    start_pdf(client, a)
    start_pdf(client, b)
    view = client.get(f"/api/documents/{b}").json()
    assert (view["state"], view["queue_position"]) == ("queued", 2)
    assert client.get(f"/api/documents/{a}").json()["queue_position"] == 1


def test_extraction_without_detection_backend_fails_the_document(sample_pdf, monkeypatch):
    import app.main as main

    class ReadOnly:
        def read_region(self, image):
            from app.pipeline.ocr.base import OcrResult
            return OcrResult(text="", confidence=0.0)
    monkeypatch.setattr("app.main._BACKEND", ReadOnly())
    sid = save_pdf(client, sample_pdf)["session_id"]
    start_pdf(client, sid)
    main.runtime().worker.run_once()
    doc = client.get(f"/api/documents/{sid}").json()
    assert (doc["state"], doc["error_type"]) == ("failed", "RuntimeError")
    assert client.post(f"/api/documents/{sid}/open").status_code == 409


def test_extraction_cancels_when_event_set(sample_pdf, stub_backend, tmp_path):
    import threading
    from app.batch.worker import Cancelled
    from app.pipeline.extract import extract

    cancel = threading.Event()
    cancel.set()  # already cancelled before the first progress emit
    seen = []

    def progress(step, detail="", current=None, total=None):
        seen.append(step)
        if cancel.is_set():
            raise Cancelled()

    with pytest.raises(Cancelled):
        extract(sample_pdf, work_dir=tmp_path, dpi=300,
                backend=stub_backend, progress=progress)
    # it unwound at the very first emit ("render"), never reaching "place"
    assert seen == ["render"]
    assert "place" not in seen


def test_upload_returns_metadata_without_extracting(sample_pdf, stub_backend):
    import app.main as main
    meta = save_pdf(client, sample_pdf)
    assert meta["fileName"] == "sample.pdf" and meta["pages"] >= 1
    work = main.runtime().data.doc_dir(meta["session_id"])
    assert (work / "input.pdf").is_file()
    assert not (work / "page.png").exists()
    assert client.get(f"/api/documents/{meta['session_id']}").json()["state"] == "staged"


@pytest.mark.parametrize("method,route", [
    ("post", "open"), ("post", "cancel"), ("post", "retry")])
def test_document_routes_reject_a_malformed_id(method, route):
    assert getattr(client, method)(f"/api/documents/not-a-valid-uuid/{route}").status_code == 404


def test_an_unknown_document_is_404():
    assert client.get("/api/documents/" + "0" * 32).status_code == 404


def test_cancel_and_retry_a_queued_drawing(sample_pdf, stub_backend):
    import app.main as main
    sid = save_pdf(client, sample_pdf)["session_id"]
    start_pdf(client, sid)
    assert client.post(f"/api/documents/{sid}/cancel").json()["state"] == "cancelled"
    assert main.runtime().worker.run_once() is None
    assert client.post(f"/api/documents/{sid}/retry").json()["state"] == "queued"
    assert main.runtime().worker.run_once() == sid


def test_read_region_returns_parsed_characteristic(sample_pdf, stub_backend):
    up = upload_pdf(client, sample_pdf)
    r = client.post("/api/read_region",
                    json={"session_id": up["session_id"], "box": [40, 40, 200, 90]})
    assert r.status_code == 200
    row = r.json()
    assert row["source"] == "manual"
    assert row["nominal"] == "1,2"
    assert row["target_region"] == [40, 40, 200, 90]
    assert row["balloon_xy"] is not None


def finished(test_client, sample_pdf, rows_fn=lambda rows: rows):
    """Extract, then seal revision 1 with `rows_fn(rows)` as the final rows."""
    up = upload_pdf(test_client, sample_pdf)
    rows = rows_fn(up["rows"])
    r = test_client.post(f"/api/session/{up['session_id']}/seal", json={
        "writer": up["writer"], "final_seq": 0, "rows": rows,
        "reviewed_ids": []})
    assert r.status_code == 200, r.text
    return up["session_id"]


def test_export_xlsx_from_the_sealed_revision(sample_pdf, stub_backend):
    sid = finished(client, sample_pdf)
    r = client.get(f"/api/documents/{sid}/export/xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="sample.xlsx"' in r.headers["content-disposition"]
    assert client.get(f"/api/documents/{sid}").json()["exported_revision"] == 1


def test_export_pdf_from_the_sealed_revision(sample_pdf, stub_backend):
    sid = finished(client, sample_pdf)
    r = client.get(f"/api/documents/{sid}/export/pdf?rev=1")
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    assert 'filename="sample_ballooned.pdf"' in r.headers["content-disposition"]


def test_an_unfinished_drawing_cannot_be_exported(sample_pdf, stub_backend):
    sid = upload_pdf(client, sample_pdf)["session_id"]
    assert client.get(f"/api/documents/{sid}/export/xlsx").status_code == 409


def test_export_rejects_an_unknown_kind_and_a_malformed_id(sample_pdf, stub_backend):
    sid = finished(client, sample_pdf)
    assert client.get(f"/api/documents/{sid}/export/docx").status_code == 404
    assert client.get("/api/documents/not-a-valid-uuid/export/xlsx").status_code == 404


def test_export_of_an_unknown_well_formed_id_creates_no_directory():
    """A well-formed but unknown id must 404 on the ROW check before any path
    under docs/ is built -- never an orphan directory for an id nobody owns."""
    import app.main as main
    doc_id = "0" * 32
    assert client.get(f"/api/documents/{doc_id}/export/xlsx").status_code == 404
    assert not main.runtime().data.doc_dir(doc_id).exists()


def test_export_with_missing_result_file_is_a_clean_409(sample_pdf, stub_backend):
    """The xlsx export reads result.json for notes/marks/title_block; if it is
    gone (same class of fault as /open's missing result.json) this is a clean
    409, not a 500 whose traceback would log the full path."""
    import app.main as main
    sid = finished(client, sample_pdf)
    (main.runtime().data.doc_dir(sid) / "result.json").unlink()
    r = client.get(f"/api/documents/{sid}/export/xlsx")
    assert r.status_code == 409
    assert sid not in r.json()["detail"]


def test_export_with_a_deleted_doc_dir_is_a_clean_409_and_stays_deleted(
        sample_pdf, stub_backend):
    """The row can outlive its files (a delete race, and routinely after a
    future retention purge): docs/<id>/ itself gone must be the same class of
    fault as any other missing input, never a 500 from the bare
    `mkdir(exist_ok=True)` finding its parent absent -- and the directory must
    not be silently recreated just because an export was attempted."""
    import shutil
    import app.main as main
    sid = finished(client, sample_pdf)
    work = main.runtime().data.doc_dir(sid)
    shutil.rmtree(work)
    r = client.get(f"/api/documents/{sid}/export/xlsx")
    assert r.status_code == 409
    assert sid not in r.json()["detail"]
    assert not work.exists()


def test_a_repeated_export_is_byte_identical(sample_pdf, stub_backend):
    """Rendered once per revision and cached, so the single export and B4's
    mass export of the same revision are the same bytes."""
    sid = finished(client, sample_pdf)
    a = client.get(f"/api/documents/{sid}/export/xlsx").content
    b = client.get(f"/api/documents/{sid}/export/xlsx").content
    assert a == b


def test_export_carries_the_title_block(sample_pdf, monkeypatch, tmp_path):
    from openpyxl import load_workbook
    from app.models import Characteristic, ExtractionResult, TitleField
    import app.main as main
    monkeypatch.setattr(main, "extract", lambda *a, **kw: ExtractionResult(
        characteristics=[Characteristic(pos=1, id="a", nominal="20")],
        title_block=[TitleField(label="Sheet / Blatt", label_en="Sheet",
                                label_de="Blatt", value="1/1")]))
    sid = finished(client, sample_pdf)
    out = tmp_path / "got.xlsx"
    out.write_bytes(client.get(f"/api/documents/{sid}/export/xlsx").content)
    wb = load_workbook(out)
    assert wb["Title Block"].cell(2, 3).value == "1/1"


@pytest.mark.parametrize("kind", ["xlsx", "pdf"])
def test_exports_drop_unconfirmed_suggestions_server_side(
        kind, sample_pdf, stub_backend, monkeypatch):
    """Defence in depth: a suggested row in the sealed revision is still not
    exported. An unconfirmed value in the client's Excel or ballooned PDF is
    exactly the failure the phantom drops exist to prevent."""
    import app.main as main
    seen = {}

    def fake_xlsx(rows, out, **kw):
        seen["rows"] = list(rows)
        out.write_bytes(b"x")

    def fake_pdf(src, rows, dpi, out_path):
        seen["rows"] = list(rows)
        out_path.write_bytes(b"%PDF")

    import app.batch.export_render as export_render
    monkeypatch.setattr(export_render, "write_workbook", fake_xlsx)
    monkeypatch.setattr(export_render, "render_ballooned_pdf", fake_pdf)
    sid = finished(client, sample_pdf,
                   lambda rows: rows + [{**rows[0], "id": "s", "suggested": True}])
    assert client.get(f"/api/documents/{sid}/export/{kind}").status_code == 200
    assert seen["rows"] and all(not r.suggested for r in seen["rows"])


def test_content_disposition_sanitises_a_hostile_filename():
    """filename comes from the untrusted upload name. CR/LF must never reach a
    header (injection), a quote must not close the quoted-string early, a
    separator has no business in a filename, and any other control character
    (tab, 0x01, DEL) is just as capable of confusing a header parser; the
    RFC 5987 form carries the real (sanitised) name for everyone else."""
    import app.main as main
    header = main._content_disposition('a"b\r\nEvil: 1/../héllo\t\x01\x7f.pdf')
    assert "\r" not in header and "\n" not in header
    assert "\t" not in header and "\x01" not in header and "\x7f" not in header
    assert header.count('filename="') == 1
    ascii_part = header.split('filename="')[1].split('"')[0]
    assert '"' not in ascii_part and "\r" not in ascii_part and "\n" not in ascii_part
    assert "/" not in ascii_part and "\\" not in ascii_part
    assert "\t" not in ascii_part and "\x01" not in ascii_part and "\x7f" not in ascii_part
    assert "filename*=UTF-8''" in header


def test_export_rev_zero_and_negative_are_404(sample_pdf, stub_backend):
    """`?rev=0` must not be treated as "no rev given" (0 is falsy, but a real
    revision argument) -- it must look up revision 0, which never exists, and
    404 like any other unknown revision."""
    sid = finished(client, sample_pdf)
    assert client.get(f"/api/documents/{sid}/export/xlsx?rev=0").status_code == 404
    assert client.get(f"/api/documents/{sid}/export/xlsx?rev=-1").status_code == 404


def test_export_pdf_with_missing_input_file_is_a_clean_409(sample_pdf, stub_backend):
    """input.pdf gone is the same class of fault as a missing result.json: a
    clean 409, never a 500 from PyMuPDF's own FileNotFoundError (a RuntimeError
    subclass, not an OSError -- it would otherwise escape uncaught)."""
    import app.main as main
    sid = finished(client, sample_pdf)
    (main.runtime().data.doc_dir(sid) / "input.pdf").unlink()
    r = client.get(f"/api/documents/{sid}/export/pdf")
    assert r.status_code == 409
    assert sid not in r.json()["detail"]


def test_export_pdf_with_corrupt_input_file_is_a_clean_409(sample_pdf, stub_backend):
    """A present but unreadable input.pdf raises PyMuPDF's FileDataError, also
    a RuntimeError subclass -- still an input fault, not a server bug."""
    import app.main as main
    sid = finished(client, sample_pdf)
    (main.runtime().data.doc_dir(sid) / "input.pdf").write_bytes(b"not a pdf at all")
    r = client.get(f"/api/documents/{sid}/export/pdf")
    assert r.status_code == 409
    assert sid not in r.json()["detail"]


def test_an_unexpected_render_fault_is_a_500_with_a_fixed_body(
        sample_pdf, stub_backend, monkeypatch):
    """A bug in the exporter itself (not a missing/unreadable input) must
    never be reported as "the files are gone" -- and the body is a fixed
    sentence, never the exception's own text."""
    import app.main as main

    def broken(*a, **kw):
        raise TypeError("boom: /some/sensitive/path")

    import app.batch.export_render as export_render
    monkeypatch.setattr(export_render, "write_workbook", broken)
    sid = finished(client, sample_pdf)
    r = client.get(f"/api/documents/{sid}/export/xlsx")
    assert r.status_code == 500
    assert r.json() == {"detail": "export failed"}


def test_mass_export_builds_a_zip_and_reports_progress(sample_pdf, stub_backend):
    import app.main as main
    sid = finished(client, sample_pdf)
    r = client.post("/api/exports", json={"document_ids": [sid]})
    assert r.status_code == 200, r.text
    job_id = r.json()["id"]
    main.runtime().exports.status(job_id)   # exists immediately, even before the thread finishes
    import time
    for _ in range(50):
        status = client.get(f"/api/exports/{job_id}").json()
        if status["state"] != "running":
            break
        time.sleep(0.05)
    assert status["state"] == "done"
    dl = client.get(f"/api/exports/{job_id}/download")
    assert dl.status_code == 200
    assert dl.headers["content-type"] == "application/zip"


def test_mass_export_of_an_unknown_id_404s():
    assert client.get(f"/api/exports/{'a' * 32}").status_code == 404


def test_downloading_an_export_still_building_is_409():
    import app.main as main
    job = main.runtime().exports.repo.create_export(["x" * 32])
    r = client.get(f"/api/exports/{job['id']}/download")
    assert r.status_code == 409


def test_mass_export_request_rejects_a_malformed_id():
    r = client.post("/api/exports", json={"document_ids": ["not-hex"]})
    assert r.status_code == 422


def test_mass_export_request_rejects_an_empty_list():
    r = client.post("/api/exports", json={"document_ids": []})
    assert r.status_code == 422


def test_mass_export_request_rejects_more_ids_than_the_cap():
    import app.main as main
    over = [f"{i:032x}" for i in range(main.MAX_EXPORT_IDS + 1)]
    r = client.post("/api/exports", json={"document_ids": over})
    assert r.status_code == 422


def test_mass_export_deduplicates_ids_preserving_order(sample_pdf, stub_backend):
    import app.main as main
    import time
    sid = finished(client, sample_pdf)
    r = client.post("/api/exports", json={"document_ids": [sid, sid]})
    assert r.status_code == 200, r.text
    job_id = r.json()["id"]
    job = main.runtime().exports.repo.get_export(job_id)
    assert job["total"] == 1
    # Wait for the job's own thread to finish before the test ends: the
    # per-test teardown closes the DB connection out from under it otherwise
    # (the same wait test_mass_export_builds_a_zip_and_reports_progress does).
    for _ in range(50):
        if main.runtime().exports.repo.get_export(job_id)["state"] != "running":
            break
        time.sleep(0.05)


def test_downloading_a_failed_export_is_409_not_500():
    """The server did not fault on THIS request -- the job's own outcome is
    a state conflict (already visible as `error_type` in the status), not a
    server error on the download call itself."""
    import app.main as main
    job = main.runtime().exports.repo.create_export(["x" * 32])
    main.runtime().exports.repo.finish_export(job["id"], "failed", error_type="boom")
    r = client.get(f"/api/exports/{job['id']}/download")
    assert r.status_code == 409
    assert r.json() == {"detail": "export failed — start it again"}


def test_downloading_a_done_export_whose_zip_was_swept_is_410():
    import app.main as main
    rt = main.runtime()
    job = rt.exports.repo.create_export(["x" * 32])
    rt.exports.repo.finish_export(job["id"], "done")
    rt.data.export_zip_path(job["id"]).unlink(missing_ok=True)
    r = client.get(f"/api/exports/{job['id']}/download")
    assert r.status_code == 410


def test_render_export_renders_once_under_concurrent_first_calls(
        sample_pdf, stub_backend, monkeypatch):
    """Two callers racing the FIRST export of a revision must not share a
    fixed temp name: the loser must not see a FileNotFoundError from a rename
    the winner already completed, or a half-written file. The lock makes the
    two calls mutually exclusive, so the exporter is expected to run exactly
    ONCE -- the second caller, arriving while the first still holds the lock
    (a `started` event lets it know that, rather than guessing with a sleep),
    must simply wait and then get the already-rendered, byte-identical file."""
    import threading
    import app.main as main

    sid = finished(client, sample_pdf)
    work = main.runtime().data.doc_dir(sid)
    sealed = main.runtime().store.revision(sid, 1)
    calls = []
    started = threading.Event()
    release = threading.Event()

    def slow_write(rows, out, **kw):
        calls.append(1)
        started.set()
        assert release.wait(timeout=5), "second caller never reached the lock"
        with open(out, "wb") as f:
            f.write(b"x" * 2000)

    import app.batch.export_render as export_render
    monkeypatch.setattr(export_render, "write_workbook", slow_write)
    results = []

    def go():
        try:
            p = main._render_export(sid, work, 1, "xlsx", sealed)
            results.append(("ok", p.read_bytes()))
        except Exception as e:
            results.append(("err", e))

    first = threading.Thread(target=go)
    first.start()
    assert started.wait(timeout=5), "first caller never started rendering"
    second = threading.Thread(target=go)
    second.start()
    # No signal for "blocked on the lock" exists (that is the point: from the
    # outside it looks identical to "not started yet"), so this is the one
    # deliberate sleep in the suite -- long enough that a second caller which
    # was NOT blocked (the bug) would have already raced past the lock, run
    # its own render, and corrupted `calls`/the output file before it is
    # released below.
    import time
    time.sleep(0.2)
    release.set()
    first.join(timeout=5)
    second.join(timeout=5)

    assert [r[0] for r in results] == ["ok", "ok"]
    assert results[0][1] == results[1][1] == b"x" * 2000
    assert len(calls) == 1
    assert list((work / "exports").glob("*.tmp.*")) == []


def test_image_rejects_bad_session_id():
    assert client.get("/api/image/not-a-valid-uuid").status_code == 404


def test_image_missing_session_returns_404():
    assert client.get("/api/image/" + "0" * 32).status_code == 404


def test_read_region_sets_needs_review_on_empty_read(monkeypatch):
    import app.main as main
    from PIL import Image

    # empty-text stub backend (StubVLMBackend.read_region returns text="")
    monkeypatch.setattr("app.main._BACKEND", StubVLMBackend(text=""))

    # a session with a page image present (read_region requires work/page.png)
    # session id must be 32 lowercase hex chars to pass _doc_dir validation
    session = "5e5537e7ae00000000000000000000ae"
    work = main._doc_dir(session)
    work.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (200, 200), "white").save(work / "page.png")

    resp = client.post("/api/read_region",
                       json={"session_id": session, "box": [10, 10, 80, 40]})
    assert resp.status_code == 200
    row = resp.json()
    assert row["needs_review"] is True
    assert row["review_reasons"] == ["empty read"]


def test_read_region_applies_active_flag_rules_like_the_pipeline(monkeypatch):
    """extract() runs _apply_active_policy after review_flags (app/pipeline/
    extract.py); a manual read through /api/read_region skipped that step
    entirely, so the same row read manually vs auto-ballooned disagreed on
    whether it needs review. diameter_sign fires on raw_text alone (no kind
    needed), so it exercises the wiring independent of ACTIVE_FLAG_RULES'
    current membership."""
    import app.pipeline.policy_rules as pr
    monkeypatch.setattr(pr, "ACTIVE_FLAG_RULES", ("diameter_sign",))

    import app.main as main
    from PIL import Image

    monkeypatch.setattr("app.main._BACKEND", StubVLMBackend(text="Ø20"))

    session = "5e5537e7ae00000000000000000000af"
    work = main._doc_dir(session)
    work.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (200, 200), "white").save(work / "page.png")

    resp = client.post("/api/read_region",
                       json={"session_id": session, "box": [10, 10, 80, 40]})
    assert resp.status_code == 200
    row = resp.json()
    assert row["needs_review"] is True
    assert pr.FLAG_REASONS["diameter_sign"] in row["review_reasons"]


def test_upload_returns_notes_field(monkeypatch, sample_pdf):
    """The upload endpoint now returns {rows, notes}; notes may be null."""
    from fastapi.testclient import TestClient
    import app.main as main
    from app.models import Characteristic, ExtractionResult, NoteBlock, Note

    monkeypatch.setattr(main, "extract", lambda *a, **kw: ExtractionResult(
        characteristics=[Characteristic(pos=1, char_type="Distance", nominal="1,2")],
        notes=NoteBlock(region=(0, 0, 10, 10),
                        notes=[Note(pos=101, text_en="A", text_de="B"),
                               Note(pos=1, parent_pos=101, sub_index=1,
                                    text_en="A1", text_de="A1")])
    ))
    test_client = TestClient(main.app)
    data = upload_pdf(test_client, sample_pdf, filename="x.pdf")
    assert "rows" in data and len(data["rows"]) == 1
    assert "notes" in data and data["notes"] is not None
    note_positions = [n["pos"] for n in data["notes"]["notes"]]
    assert note_positions == [101, 1]


def test_upload_returns_null_notes_when_extract_returns_none(monkeypatch, sample_pdf):
    from fastapi.testclient import TestClient
    import app.main as main
    from app.models import Characteristic, ExtractionResult

    monkeypatch.setattr(main, "extract", lambda *a, **kw: ExtractionResult(
        characteristics=[Characteristic(pos=1)], notes=None))
    test_client = TestClient(main.app)
    data = upload_pdf(test_client, sample_pdf, filename="x.pdf")
    assert data["notes"] is None


def test_delete_document_removes_its_files_and_row(sample_pdf):
    import app.main as main
    sid = save_pdf(client, sample_pdf)["session_id"]
    work = main.runtime().data.doc_dir(sid)
    assert work.is_dir()
    assert client.delete(f"/api/documents/{sid}").status_code == 200
    assert not work.exists()
    assert client.get(f"/api/documents/{sid}").status_code == 404


def test_delete_refuses_the_drawing_the_worker_is_still_running(
        sample_pdf, stub_backend, monkeypatch):
    """Cancel is cooperative: the row says `cancelled` at once, but the run
    goes on until its next progress report. The repo allows deleting a
    cancelled row, so the API must refuse the worker's CURRENT job itself --
    otherwise the files would be removed under a run that is still writing.
    The fake extract cancels and tries to delete from inside the run, which
    is exactly that window, with no threads or timing involved."""
    import app.main as main
    seen = {}

    def extract_that_is_cancelled_midway(*a, **k):
        sid = seen["sid"]
        assert client.post(f"/api/documents/{sid}/cancel").json()["state"] == "cancelled"
        r = client.delete(f"/api/documents/{sid}")
        seen["status"], seen["body"] = r.status_code, r.json()
        seen["dir_kept"] = main.runtime().data.doc_dir(sid).is_dir()
        raise RuntimeError("never reached a result")
    monkeypatch.setattr(main, "extract", extract_that_is_cancelled_midway)
    sid = seen["sid"] = save_pdf(client, sample_pdf)["session_id"]
    start_pdf(client, sid)
    assert main.runtime().worker.run_once() == sid
    assert (seen["status"], seen["dir_kept"]) == (409, True), seen["body"]
    assert sid not in seen["body"]["detail"]
    # Once the run has returned, the cancelled drawing can be removed.
    assert client.get(f"/api/documents/{sid}").json()["state"] == "cancelled"
    assert client.delete(f"/api/documents/{sid}").status_code == 200


def test_delete_missing_document_is_idempotent():
    assert client.delete("/api/documents/" + "0" * 32).status_code == 200


def test_delete_rejects_bad_document_id():
    assert client.delete("/api/documents/not-a-valid-uuid").status_code == 404


@pytest.mark.parametrize("path", ["/", "/js/main.js", "/js/api.js"])
def test_static_assets_force_revalidation(path):
    """Static HTML/JS must carry Cache-Control: no-cache so a rebuilt backend
    is never paired with a stale cached frontend (the source of the
    `payload.rows.map` crash after a container rebuild)."""
    r = client.get(path)
    assert r.status_code == 200, r.text
    assert "no-cache" in r.headers.get("cache-control", "")


def test_upload_returns_marks_field(monkeypatch, sample_pdf):
    """Extraction result carries `marks` alongside `notes`."""
    from fastapi.testclient import TestClient
    import app.main as main
    from app.models import (
        Characteristic, ExtractionResult, MarkBlock, Mark,
    )

    monkeypatch.setattr(main, "extract", lambda *a, **kw: ExtractionResult(
        characteristics=[Characteristic(pos=1, char_type="Distance", nominal="1,2")],
        marks=MarkBlock(region=(0, 0, 10, 10),
                        marks=[Mark(pos=101, text_en="A", text_de="B"),
                               Mark(pos=102, text_en="C", text_de="D")])
    ))
    test_client = TestClient(main.app)
    data = upload_pdf(test_client, sample_pdf, filename="x.pdf")
    assert "marks" in data and data["marks"] is not None
    mark_positions = [m["pos"] for m in data["marks"]["marks"]]
    assert mark_positions == [101, 102]


def test_upload_returns_null_marks_when_extract_returns_none(monkeypatch, sample_pdf):
    from fastapi.testclient import TestClient
    import app.main as main
    from app.models import Characteristic, ExtractionResult

    monkeypatch.setattr(main, "extract", lambda *a, **kw: ExtractionResult(
        characteristics=[Characteristic(pos=1)], marks=None))
    test_client = TestClient(main.app)
    data = upload_pdf(test_client, sample_pdf, filename="x.pdf")
    assert data["marks"] is None


# --- the suggestion tray (docs/plans/2026-10-07-suggestion-tray-design.md) ---

def test_extract_payload_merges_suggestions_flagged(monkeypatch, sample_pdf):
    from fastapi.testclient import TestClient
    import app.main as main
    from app.models import Characteristic, ExtractionResult
    monkeypatch.setattr(main, "extract", lambda *a, **kw: ExtractionResult(
        characteristics=[Characteristic(pos=1, id="a", nominal="20")],
        suggestions=[Characteristic(pos=0, id="b", nominal="7",
                                    suggested=True)]))
    data = upload_pdf(TestClient(main.app), sample_pdf, filename="x.pdf")
    assert [(r["id"], r["suggested"]) for r in data["rows"]] == \
        [("a", False), ("b", True)]


# --- the queue API: failure paths, lifespan (batch design §1, §3) ---

@pytest.mark.parametrize("damage", ["missing", "empty_object", "not_json"])
def test_open_without_its_result_file_is_a_clean_409(damage, sample_pdf, stub_backend):
    """A ready drawing whose result.json is gone or damaged must not become a
    500 whose traceback logs the document's full path; the body says nothing
    about the file either."""
    import app.main as main
    sid = upload_pdf(client, sample_pdf)["session_id"]
    result = main.runtime().data.doc_dir(sid) / "result.json"
    if damage == "missing":
        result.unlink()
    elif damage == "empty_object":
        result.write_text("{}")
    else:
        result.write_text("{not json")
    r = TestClient(main.app, raise_server_exceptions=False).post(f"/api/documents/{sid}/open")
    assert r.status_code == 409
    assert r.json() == {"detail": "drawing files unavailable"}


def test_a_failed_open_does_not_rotate_the_writer(sample_pdf, stub_backend):
    """Opening rotates the writer token, which locks out the tab that held it.
    An open refused for missing files must not do that for nothing."""
    import app.main as main
    up = upload_pdf(client, sample_pdf)
    sid, writer = up["session_id"], up["writer"]
    (main.runtime().data.doc_dir(sid) / "result.json").unlink()
    assert client.post(f"/api/documents/{sid}/open").status_code == 409
    r = client.post(f"/api/session/{sid}/events", json={"writer": writer, "events": [
        {"seq": 1, "type": "accept", "ids": []}]})
    assert r.status_code == 200, r.text


def test_retry_refuses_the_drawing_the_worker_is_still_running(
        sample_pdf, stub_backend, monkeypatch):
    """A cancelled row can still be the worker's current run (cancel is
    cooperative). Retrying it then would requeue it while the old run is
    still unwinding -- and that run's settle would silently undo the retry or
    the cancel. Refused until the run has returned."""
    import app.main as main
    seen = {}

    def extract_that_is_cancelled_midway(*a, **k):
        sid = seen["sid"]
        client.post(f"/api/documents/{sid}/cancel")
        r = client.post(f"/api/documents/{sid}/retry")
        seen["status"], seen["body"] = r.status_code, r.json()
        raise RuntimeError("never reached a result")
    monkeypatch.setattr(main, "extract", extract_that_is_cancelled_midway)
    sid = seen["sid"] = save_pdf(client, sample_pdf)["session_id"]
    start_pdf(client, sid)
    assert main.runtime().worker.run_once() == sid
    assert seen["status"] == 409, seen["body"]
    assert client.get(f"/api/documents/{sid}").json()["state"] == "cancelled"
    assert client.post(f"/api/documents/{sid}/retry").json()["state"] == "queued"


def test_a_running_drawing_reports_its_progress(sample_pdf, monkeypatch):
    import app.main as main
    seen = {}

    def extract_reporting(pdf_path, work_dir, dpi, backend, progress):
        progress("detect", "tile", 1, 3)
        seen["doc"] = client.get(f"/api/documents/{seen['sid']}").json()
        raise RuntimeError("stop here")
    monkeypatch.setattr(main, "extract", extract_reporting)
    sid = seen["sid"] = save_pdf(client, sample_pdf)["session_id"]
    start_pdf(client, sid)
    main.runtime().worker.run_once()
    assert seen["doc"]["state"] == "extracting"
    assert seen["doc"]["progress"] == {"step": "detect", "detail": "tile",
                                       "current": 1, "total": 3}
    assert client.get(f"/api/documents/{sid}").json()["progress"] is None


def test_the_lifespan_starts_and_stops_the_worker():
    import app.main as main
    with TestClient(main.app) as c:
        assert c.get("/api/health").json()["worker_running"] is True
    assert main._RUNTIME is None


def test_after_shutdown_requests_are_refused_not_served_by_a_dead_worker():
    """A runtime built after shutdown would accept and queue drawings that no
    worker ever extracts. Refuse instead; the next lifespan start clears it."""
    import app.main as main
    with TestClient(main.app):
        pass
    r = TestClient(main.app, raise_server_exceptions=False).get("/api/documents/" + "0" * 32)
    assert r.status_code == 503
    assert main._RUNTIME is None
    with TestClient(main.app) as c:
        assert c.get("/api/health").json()["worker_running"] is True


def test_an_unwritable_review_dir_fails_the_boot(tmp_path, monkeypatch):
    import os
    import app.main as main
    if os.geteuid() == 0:
        pytest.skip("root bypasses permission bits")
    d = tmp_path / "ro-reviews"
    d.mkdir()
    os.chmod(d, 0o500)
    monkeypatch.setenv("SINDRI_REVIEW_DIR", str(d))
    try:
        with pytest.raises(OSError):
            with TestClient(main.app):
                pass
    finally:
        os.chmod(d, 0o700)
    assert main._RUNTIME is None


def test_a_failed_worker_start_leaves_no_runtime_behind(monkeypatch):
    import app.main as main
    from app.batch.worker import Worker

    def refuse(self):
        raise RuntimeError("another process holds the worker")
    monkeypatch.setattr(Worker, "start", refuse)
    with pytest.raises(RuntimeError):
        with TestClient(main.app):
            pass
    assert main._RUNTIME is None
